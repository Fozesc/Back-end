"""
E-mail do backup: a planilha sai com uma aba por assunto (titulos, clientes, vales...) e sem
auditoria nem senha de usuario; o relatorio PDF e a planilha saem com a senha; a copia do banco
so vai para quem foi marcado na aba Backups; cada arquivo sai na sua frequencia; e o "desde o
ultimo relatorio" so anda quando o envio da certo.
Banco SEPARADO (fozesc_teste_backup_email), criado e apagado pelo teste. Nao manda e-mail:
o Gmail e' trocado por um falso.

Rode:  ./venv/bin/python test_backup_email.py
"""
import os
import re
import subprocess
import tempfile
from datetime import date, datetime, timedelta

URL_REAL = os.getenv('DATABASE_URL') or ''
if not URL_REAL:
    import dotenv
    dotenv.load_dotenv()
    URL_REAL = os.environ['DATABASE_URL']

BANCO_TESTE = 'fozesc_teste_backup_email'
URL_TESTE = re.sub(r'/[^/]+$', f'/{BANCO_TESTE}', URL_REAL)
os.environ['DATABASE_URL'] = URL_TESTE
os.environ.setdefault('JWT_SECRET_KEY', 'teste')
for k in ('POSTGRES_USER', 'POSTGRES_PASSWORD', 'POSTGRES_DB'):
    os.environ.setdefault(k, 'nao-usado')   # o gerar_backup exige; a conexao vem da DATABASE_URL

import psycopg2
from psycopg2.extensions import ISOLATION_LEVEL_AUTOCOMMIT

_app = _db = None


def sql_admin(*comandos):
    conn = psycopg2.connect(URL_REAL)
    conn.set_isolation_level(ISOLATION_LEVEL_AUTOCOMMIT)
    cur = conn.cursor()
    for c in comandos:
        cur.execute(c)
    cur.close()
    conn.close()


def main():
    global _app, _db
    sql_admin(f'DROP DATABASE IF EXISTS {BANCO_TESTE}', f'CREATE DATABASE {BANCO_TESTE}')
    try:
        from app import create_app, db
        from app.models.domain import (AuditLog, BackupConfig, Check, CheckExtension, Client, ClientNote,
                                       CompanySettings, Evento, Operation, Recorrente, Transaction, User, Vale)
        from openpyxl import load_workbook
        import gerar_backup as g
        import relatorio_backup

        app = create_app()
        _app, _db = app, db
        hoje = date.today()
        with app.app_context():
            db.session.add_all([CompanySettings(company_name='Fozesc Teste', cnpj='12.345.678/0001-90'),
                                User(name='Ana', email='ana@x.com', password_hash='HASH-QUE-NAO-PODE-SAIR', role='Admin'),
                                AuditLog(user_name='Ana', action='LOGIN', target='Auth', description='AUDITORIA-NAO-SAI')])
            maria = Client(name='Maria Júlia', credit_limit=5000)
            db.session.add(maria)
            db.session.flush()
            op = Operation(client_id=maria.id, operation_date=hoje - timedelta(days=3), total_face_value=1000,
                           total_interest=40, iof_amount=4, total_net_value=956)
            db.session.add(op)
            db.session.flush()
            atrasado = Check(operation_id=op.id, issuer_name='João “Zé” Silva', number='9', amount=600,
                             due_date=hoje - timedelta(days=2))
            db.session.add_all([atrasado, Check(operation_id=op.id, issuer_name='Ana', amount=400, due_date=hoje,
                                                status='Pago', payment_date=hoje, paid_amount=400)])
            db.session.flush()
            vale = Vale(descricao='Adiantamento — Pedro', valor=300, data=hoje - timedelta(days=1), conta='BB')
            db.session.add(vale)
            db.session.flush()
            db.session.add_all([
                Transaction(date=hoje - timedelta(days=3), amount=-956, type='saida', origin='BB', category='Compra de Ativos',
                            description=f'Pgto Borderô #{op.id}', operation_id=op.id),
                Transaction(date=hoje - timedelta(days=3), amount=0, valor_informativo=40, type='entrada', origin='BB',
                            category='Informativo', description='INFORMATIVO-NAO-SAI'),
                Transaction(date=hoje, amount=400, type='entrada', origin='Dinheiro', category='Recebimento de Cheque',
                            description='Recebimento Cheque #10'),
                Transaction(date=hoje - timedelta(days=1), amount=-300, type='saida', origin='BB', category='Vale',
                            description='Vale', vale_id=vale.id),
                Transaction(date=hoje, amount=100, type='entrada', origin='BB', category='Vale',
                            description='Pagamento parcial do vale', vale_id=vale.id),
                CheckExtension(check_id=atrasado.id, prorrogation_date=datetime.now(), old_due_date=hoje - timedelta(days=9),
                               new_due_date=hoje - timedelta(days=2), days_added=7, fee_amount=30),
                Recorrente(descricao='Aluguel', valor=1500, proximo=hoje + timedelta(days=5), dia=1),
                Evento(data=hoje, titulo='Ligar para a Maria', autor='Ana'),
                ClientNote(client_id=maria.id, autor='Ana', texto='Paga sempre em dia'),
            ])
            db.session.commit()

        # ------------------------------------------------------------ planilha
        pasta = tempfile.mkdtemp()
        g.exportar_para_xlsx(pasta, 'TESTE')
        livro = load_workbook(os.path.join(pasta, 'TESTE.xlsx'))
        assert livro.sheetnames == list(g.ABAS), livro.sheetnames
        assert not any('audit' in n.lower() or 'usu' in n.lower() for n in livro.sheetnames)
        tudo = ' '.join(str(c.value) for f in livro.worksheets for linha in f.iter_rows() for c in linha)
        for proibido in ('HASH-QUE-NAO-PODE-SAIR', 'AUDITORIA-NAO-SAI', 'INFORMATIVO-NAO-SAI'):
            assert proibido not in tudo, proibido
        aba = lambda n: [[c.value for c in linha] for linha in livro[n].iter_rows()]
        titulos = aba('Títulos')
        sit = titulos[0].index('Situação')
        assert sorted(l[sit] for l in titulos[1:]) == ['Atrasado', 'Pago'], 'Aguardando vencido aparece como Atrasado'
        clientes = aba('Clientes')
        assert clientes[1][clientes[0].index('Em aberto')] == 600
        vales = aba('Vales')
        assert (vales[1][vales[0].index('Pago')], vales[1][vales[0].index('Falta pagar')]) == (100, 200), vales
        assert len(aba('Caixa')) == 5 and len(aba('Prorrogações')) == 2 and len(aba('Despesas recorrentes')) == 2

        # ------------------------------------------------------------ relatorio PDF
        d = relatorio_backup.dados(hoje - timedelta(days=7), hoje)
        assert d['empresa'] == 'Fozesc Teste' and len(d['borderos']) == 1 and len(d['caixa']['linhas']) == 4
        assert relatorio_backup.titulo_do_periodo(date(2026, 10, 1), date(2026, 10, 30)) == 'Relatório de outubro/2026'
        assert relatorio_backup.titulo_do_periodo(date(2026, 10, 23), date(2026, 10, 30)) == 'Relatório de 23/10 a 30/10/2026'
        aberto = relatorio_backup.gerar(d, os.path.join(pasta, 'aberto.pdf'), 'Relatório')
        fechado = relatorio_backup.gerar(d, os.path.join(pasta, 'fechado.pdf'), 'Relatório', senha='senha-do-teste-123')
        assert open(aberto, 'rb').read().startswith(b'%PDF') and b'/Encrypt' not in open(aberto, 'rb').read()
        assert b'/Encrypt' in open(fechado, 'rb').read() and b'/AESV3' in open(fechado, 'rb').read(), 'AES-256 com senha'
        vazio = relatorio_backup.dados(hoje + timedelta(days=30), hoje + timedelta(days=31))
        relatorio_backup.gerar(vazio, os.path.join(pasta, 'vazio.pdf'), 'Relatório')   # periodo sem nada nao quebra

        # ------------------------------------------------ quem recebe o que (o enviar_backups.py e' falso aqui)
        g.ARQUIVO_ULTIMO_EMAIL = os.path.join(pasta, 'ULTIMO_EMAIL.json')
        g.DIR_RELATORIOS = os.path.join(pasta, 'relatorios')
        rodar_de_verdade = subprocess.run
        chamadas, resposta = [], {'codigo': 0}
        def envio_falso(cmd, env=None, **kw):
            chamadas.append(env)
            return subprocess.CompletedProcess(cmd, resposta['codigo'], stdout='enviado (teste)', stderr='')
        g.subprocess.run = envio_falso
        d = date.fromisoformat

        os.environ.pop('BACKUP_EMAIL', None)
        assert g.enviar_copia(d('2026-10-30')) is False and not chamadas, 'sem ninguem na lista nao envia'

        os.environ.update(BACKUP_EMAIL='env@recebe.com', BACKUP_SENHA='senha-do-env-12345')
        assert g.enviar_copia(d('2026-10-30')) is True, 'aba ainda nao salva: vale o .env'
        env = chamadas[-1]
        assert (env['BACKUP_EMAIL'], env['BACKUP_ARQUIVOS'], env['BACKUP_SENHA']) == \
            ('env@recebe.com', 'banco,planilha,relatorio', 'senha-do-env-12345'), env
        assert (env['BACKUP_RELATORIO_TITULO'], env['BACKUP_RELATORIO_PERIODO']) == \
            ('Relatório de outubro/2026', '01/10/2026 a 30/10/2026')
        assert b'/Encrypt' in open(env['BACKUP_RELATORIO'], 'rb').read()
        assert g.ultimo_envio() == d('2026-10-30')

        with app.app_context():
            db.session.add(BackupConfig(id=1, senha='senha-da-aba-12345', freq_banco='diario', freq_planilha='semanal',
                                        freq_relatorio='semanal',
                                        destinatarios=[{'email': 'lucas@x.com', 'banco': True},
                                                       {'email': 'socio@x.com', 'banco': False}]))
            db.session.commit()
        enviados = lambda: [(c['BACKUP_EMAIL'], c['BACKUP_ARQUIVOS']) for c in chamadas]

        chamadas.clear()
        assert g.enviar_copia(d('2026-11-04')) is True
        assert enviados() == [('lucas@x.com', 'banco')], 'quarta: so o banco, so para o Lucas'
        assert chamadas[0]['BACKUP_SENHA'] == 'senha-da-aba-12345', 'a senha da aba vale mais que a do .env'
        assert g.ultimo_envio() == d('2026-10-30'), 'sem relatorio o "desde o ultimo" nao anda'

        chamadas.clear()
        assert g.enviar_copia(d('2026-11-06')) is True
        assert enviados() == [('lucas@x.com', 'banco,planilha,relatorio'), ('socio@x.com', 'planilha,relatorio')]
        assert chamadas[1]['BACKUP_RELATORIO_PERIODO'] == '30/10/2026 a 06/11/2026', chamadas[1]['BACKUP_RELATORIO_PERIODO']
        assert g.ultimo_envio() == d('2026-11-06')

        resposta['codigo'] = 1
        assert g.enviar_copia(d('2026-11-13')) is False
        assert g.ultimo_envio() == d('2026-11-06'), 'falhou: o proximo relatorio ainda vem desde o que chegou'
        with app.app_context():
            ultimo = db.session.get(BackupConfig, 1).ultimo_envio
        assert ultimo['ok'] is False and 'FALHOU' in ultimo['texto'] and 'socio@x.com' in ultimo['texto'], ultimo
        resposta['codigo'] = 0
        chamadas.clear()
        g.enviar_copia(d('2026-11-19'))
        assert chamadas[-1]['BACKUP_RELATORIO_PERIODO'] == '06/11/2026 a 19/11/2026', 'sexta 20/11 e feriado: sai na quinta'
        with app.app_context():
            assert db.session.get(BackupConfig, 1).ultimo_envio['ok'] is True

        # ------------------------------------------------ botao "Enviar agora": o worker atende o pedido uma vez
        feitos = []
        g.executar_backup = lambda tipo: feitos.append(tipo)
        chamadas.clear()
        g.atender_pedido()
        assert not feitos and not chamadas, 'sem pedido nao faz nada'
        with app.app_context():
            db.session.get(BackupConfig, 1).envio_pedido_em = datetime.now()
            db.session.commit()
        g.atender_pedido()
        assert feitos == ['DIARIO'], 'backup novo antes de mandar'
        assert enviados() == [('lucas@x.com', 'banco,planilha,relatorio'), ('socio@x.com', 'planilha,relatorio')]
        g.atender_pedido()
        assert feitos == ['DIARIO'], 'o pedido e atendido uma vez so'
        with app.app_context():
            assert db.session.get(BackupConfig, 1).envio_pedido_em is None

        # ------------------------------------------------ a mensagem de verdade (sem o Gmail)
        g.subprocess.run = rodar_de_verdade
        import enviar_backups as e
        import gzip, io, sys, msoffcrypto
        mensagens = []
        class GmailFalso:
            def __init__(self, *a, **k): pass
            def __enter__(self): return self
            def __exit__(self, *a): return False
            def login(self, *a): pass
            def send_message(self, msg): mensagens.append(msg)
        e.smtplib.SMTP_SSL = GmailFalso
        e.limpar_antigos = lambda *a: 0
        dump = b'-- x\n-- PostgreSQL database dump complete\n'
        with gzip.open(os.path.join(pasta, 'FOZESC_DIARIO.sql.gz'), 'wb') as f:
            f.write(dump)
        os.environ.update(BACKUP_NOME='Fozesc', BACKUP_PASTA=pasta, BACKUP_SMTP_USER='envia@x.com',
                          BACKUP_SMTP_PASS='abcd efgh', BACKUP_SENHA='senha-do-teste-123', BACKUP_RELATORIO=fechado,
                          BACKUP_RELATORIO_TITULO='Relatório de outubro/2026',
                          BACKUP_RELATORIO_PERIODO='01/10/2026 a 30/10/2026')
        sys.argv = ['enviar_backups.py']
        hoje_br = datetime.now().strftime('%d/%m/%Y')

        os.environ.update(BACKUP_EMAIL='socio@x.com', BACKUP_ARQUIVOS='planilha,relatorio')
        e.main()
        msg = mensagens[-1]
        assert msg['Subject'] == f'Backup Fozesc {hoje_br} — Relatório de outubro/2026 e planilha', msg['Subject']
        anexos = {a.get_filename(): a.get_payload(decode=True) for a in msg.iter_attachments()}
        planilha = f"Fozesc - Planilha {datetime.now():%d-%m-%Y}.xlsx"
        assert sorted(anexos) == sorted(['Fozesc - Relatório de outubro-2026.pdf', planilha]), sorted(anexos)
        corpo = msg.get_body().get_content()
        assert 'openssl' not in corpo and 'ATENÇÃO' not in corpo, corpo
        assert 'Resumo do movimento de 01/10/2026 a 30/10/2026' in corpo and 'Títulos, Clientes, Borderôs' in corpo, corpo
        cifrada = msoffcrypto.OfficeFile(io.BytesIO(anexos[planilha]))
        assert cifrada.is_encrypted(), 'a planilha so abre com a senha'
        cifrada.load_key(password='senha-do-teste-123')
        aberta = io.BytesIO()
        cifrada.decrypt(aberta)
        assert load_workbook(aberta).sheetnames == list(g.ABAS)

        os.environ.update(BACKUP_EMAIL='lucas@x.com', BACKUP_ARQUIVOS='banco,planilha,relatorio')
        e.main()
        msg = mensagens[-1]
        assert msg['Subject'].endswith('Relatório de outubro/2026, planilha e cópia do banco'), msg['Subject']
        anexos = {a.get_filename(): a.get_payload(decode=True) for a in msg.iter_attachments()}
        banco = f"fozesc-banco-{datetime.now():%Y-%m-%d}.sql.gz.enc"
        aberto = subprocess.run(['openssl', 'enc', '-d', *e.CIFRA, '-pass', 'pass:senha-do-teste-123'],
                                input=anexos[banco], capture_output=True, check=True).stdout
        assert gzip.decompress(aberto) == dump, 'a copia do banco abre com a senha e e a mesma'
        assert f'-in {banco} -out {banco[:-4]}' in msg.get_body().get_content()

        os.environ.update(BACKUP_EMAIL='socio@x.com', BACKUP_ARQUIVOS='planilha,relatorio', BACKUP_RELATORIO='')
        e.main()
        msg = mensagens[-1]
        assert msg['Subject'] == f'Backup Fozesc {hoje_br} — ATENÇÃO: Planilha', msg['Subject']
        assert 'O relatório não foi gerado.' in msg.get_body().get_content()
        print('OK: e-mail do backup - planilha com uma aba por assunto (sem auditoria, sem senha de usuário, sem linha '
              'informativa), relatório PDF e planilha com senha, cópia do banco só para quem foi marcado, cada arquivo '
              'na sua frequência, a senha da aba Backups no lugar da do .env, e "desde o último relatório" só anda '
              'quando ele chega a todos.')
    finally:
        if _app is not None:
            with _app.app_context():
                _db.engine.dispose()
        sql_admin(f'DROP DATABASE IF EXISTS {BANCO_TESTE}')


if __name__ == '__main__':
    main()
